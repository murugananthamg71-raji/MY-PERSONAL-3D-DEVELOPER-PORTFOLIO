import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_report():
    doc = docx.Document()

    # Set page margins (1 inch all around)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Palette
    COLOR_PRIMARY = RGBColor(15, 43, 89)      # Deep Navy #0F2B59
    COLOR_SECONDARY = RGBColor(2, 132, 199)   # Accent Cyan #0284C7
    COLOR_TEXT = RGBColor(31, 41, 55)         # Charcoal #1F2937
    COLOR_MUTED = RGBColor(100, 116, 139)     # Slate Gray #64748B
    HEX_HEADER = "0F2B59"
    HEX_ALT = "F8FAFC"
    HEX_BORDER = "CBD5E1"
    HEX_CALLOUT_BG = "F1F5F9"

    # Set Normal Style font
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = COLOR_TEXT
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)

    # ------------------ COVER PAGE / TITLE ------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(8)
    title_run = title_p.add_run("PROJECT REPORT & TECHNICAL SPECIFICATION")
    title_run.font.size = Pt(13)
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_SECONDARY

    main_title_p = doc.add_paragraph()
    main_title_p.paragraph_format.space_after = Pt(12)
    main_title_run = main_title_p.add_run("Personal 3D Developer Portfolio & Web Architecture")
    main_title_run.font.size = Pt(24)
    main_title_run.font.bold = True
    main_title_run.font.color.rgb = COLOR_PRIMARY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(24)
    sub_run = sub_p.add_run("Engineering Specification, WebGL 3D Integration, Full-Stack Showcase & Complete System Architecture")
    sub_run.font.size = Pt(12)
    sub_run.font.italic = True
    sub_run.font.color.rgb = COLOR_MUTED

    # Meta Table
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Candidate Name", "MURUGANANTHAM G"),
        ("Engineering Role", "Full Stack Developer & AI Enthusiast"),
        ("Education & Institution", "B.Tech Information Technology, Karpagam Institute of Technology, Coimbatore"),
        ("Contact & Location", "+91 88706 42691 • murugananthamg@gmail.com • Coimbatore, Tamil Nadu, India"),
        ("GitHub Repository", "https://github.com/murugananthamg71-raji/MY-PERSONAL-3D-DEVELOPER-PORTFOLIO"),
        ("Live Development URL", "http://localhost:3000 / Production Deployment via Vercel")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "E2E8F0")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, 100, 100, 140, 140)
        set_cell_margins(c1, 100, 100, 140, 140)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = COLOR_PRIMARY

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(10)
        if "http" in v or "@" in v:
            r1.font.bold = True
            r1.font.color.rgb = COLOR_SECONDARY

    doc.add_page_break()

    # Helper functions for sections
    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_callout(text, italic=True):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, HEX_CALLOUT_BG)
        set_cell_margins(cell, 140, 140, 180, 180)
        # Left border highlight
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="none"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="0284C7"/>
                <w:bottom w:val="none"/>
                <w:right w:val="none"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.font.size = Pt(10)
        run.font.italic = italic
        run.font.color.rgb = COLOR_TEXT
        doc.add_paragraph() # spacing

    # ------------------ 1. EXECUTIVE SUMMARY ------------------
    add_h1("1. Executive Summary")
    doc.add_paragraph(
        "The Personal 3D Developer Portfolio is an enterprise-grade, highly interactive personal software engineering portal engineered by Muruganantham G. "
        "Built atop the Next.js 16 App Router framework and React 19, the application unifies modern high-performance full-stack web engineering, "
        "WebGL-powered 3D computer graphics, programmatic sound effects, and clean, responsive component architecture into a unified developer showcase."
    )
    doc.add_paragraph(
        "Unlike generic static template portfolios, this system features a custom Chibi3D WebGL interactive mechanical keyboard where each keycap is mapped to a real "
        "technology in the developer's tech stack (React, Python, TypeScript, Docker, PostgreSQL, etc.), responding dynamically to mouse clicks, physical keyboard typing, "
        "and hover interactions. The application delivers multi-route direct navigation, verified educational credentials, professional internship experiences with authentic "
        "offer letter documentation, a curated showcase of 7 real-world projects, and a fully functional server-side email dispatch system."
    )

    add_callout(
        "Key Project Goal: To bridge full-stack system architecture with applied artificial intelligence, presenting Muruganantham G's technical competencies, "
        "real-world engineering prototypes, academic milestones, and professional readiness in an engaging, production-ready interface."
    )

    # ------------------ 2. CORE ARCHITECTURAL PILLARS ------------------
    add_h1("2. Core Architectural Pillars & Design Innovations")

    add_h2("2.1 Interactive WebGL 3D Mechanical Keyboard")
    doc.add_paragraph(
        "At the focal center of the landing view is an interactive 3D mechanical keyboard powered by @chibi3d/runtime and Three.js. "
        "Key capabilities include:"
    )
    p = doc.add_paragraph()
    p.add_run("• Raycast Hover & Click Detection: ").font.bold = True
    p.add_run("Each keycap corresponds to a specific programming language, framework, database, or tool with real-time emissive lighting.")
    p = doc.add_paragraph()
    p.add_run("• Physical Hardware Keyboard Binding: ").font.bold = True
    p.add_run("Typing letters on a physical computer keyboard triggers corresponding 3D key depressions and mechanical switch audio in real time.")
    p = doc.add_paragraph()
    p.add_run("• Audio Synthesis: ").font.bold = True
    p.add_run("A programmatic Web Audio API synthesizes realistic mechanical keypress audio feedback on stroke and release.")
    p = doc.add_paragraph()
    p.add_run("• Performance Adaptability & Accessibility: ").font.bold = True
    p.add_run("Automatic DPR scaling based on hardware GPU benchmarks, with complete graceful HTML fallback grids for screen-readers and low-power devices.")

    add_h2("2.2 Multi-Route Direct Page Navigation Architecture")
    doc.add_paragraph(
        "Rather than trapping visitors in an endless, sluggish single-page scroll-wheel marathon, the portfolio implements a crisp, modern multi-route architecture. "
        "The Home page serves as a focused, zero-scroll 3D landing screen, while all sections exist as dedicated routes with complete URL routing, direct link sharing, and instant loading:"
    )

    nav_table = doc.add_table(rows=10, cols=3)
    nav_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    nav_headers = ["Page / Section", "Route", "Functional Purpose"]
    for j, h in enumerate(nav_headers):
        cell = nav_table.rows[0].cells[j]
        set_cell_background(cell, HEX_HEADER)
        set_cell_margins(cell, 120, 120, 140, 140)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    routes_data = [
        ("Home", "/", "3D Interactive Keyboard, Headline, Animated Roles, Action Buttons, Marquee Ribbon"),
        ("About Me", "/about", "Developer Biography, Technical Focus Highlights, Systems Philosophy"),
        ("Skills", "/skills", "Categorized Skill Groups (Frontend, Backend, DB, Tools, AI) & Keycaps"),
        ("Projects", "/projects", "7 Production Projects with Modals, Architecture, GitHub & Live Links"),
        ("Certifications", "/certifications", "10 Official Verified Credentials with Dual-Mode View & PDF Downloads"),
        ("Experience", "/experience", "AEX Xenora, YuvaIntern, Texmo with Official Verified Offer Letters"),
        ("Education", "/education", "B.Tech IT, Diploma EEE, SSLC Academic Journey & Distinction Highlights"),
        ("Resume", "/resume", "Interactive PDF Resume Viewer, Skills Matrix & Direct PDF Download"),
        ("Contact", "/contact", "Validated Contact Form, Direct Phone (+91 88706 42691), Social Profiles")
    ]
    for i, row_data in enumerate(routes_data, start=1):
        row = nav_table.rows[i]
        bg = HEX_ALT if i % 2 == 1 else "FFFFFF"
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 90, 90, 120, 120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9)
            if j == 0:
                r.font.bold = True
            elif j == 1:
                r.font.color.rgb = COLOR_SECONDARY
                r.font.bold = True

    doc.add_paragraph() # spacing

    # ------------------ 3. TECHNICAL STACK ------------------
    add_h1("3. Comprehensive Technology Stack")
    doc.add_paragraph("The application leverages cutting-edge web technologies selected for speed, developer ergonomics, and aesthetic precision:")

    tech_table = doc.add_table(rows=8, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_headers = ["Layer", "Technology", "Architectural Role & Benefit"]
    for j, h in enumerate(tech_headers):
        cell = tech_table.rows[0].cells[j]
        set_cell_background(cell, HEX_HEADER)
        set_cell_margins(cell, 120, 120, 140, 140)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)

    tech_data = [
        ("Framework", "Next.js 16.2.2 (Turbopack)", "App Router, Server-Side Rendering (SSR), Static Optimization, Sub-second HMR"),
        ("UI Library", "React 19.2.4", "Concurrent rendering, modern hook primitives, component isolation"),
        ("Language", "TypeScript 5.x", "End-to-end static type safety across data schemas, components, and API routes"),
        ("Styling", "Tailwind CSS & Shadcn UI", "Utility-first design tokens, Radix UI accessibility primitives, dark cosmic space theme"),
        ("3D & Graphics", "Three.js & @chibi3d/runtime", "WebGL 2.0 accelerated 3D scene rendering, shadows, bloom, and raycasting"),
        ("Animation", "Framer Motion & GSAP", "Buttery smooth layout transitions, infinite marquee ticker, spring mechanics"),
        ("Backend & Email", "Next.js Route Handlers & Resend API", "Serverless email dispatch, Zod runtime schema validation, Mailto fallback")
    ]
    for i, row_data in enumerate(tech_data, start=1):
        row = tech_table.rows[i]
        bg = HEX_ALT if i % 2 == 1 else "FFFFFF"
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 90, 90, 120, 120)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9)
            if j == 0:
                r.font.bold = True
            elif j == 1:
                r.font.bold = True
                r.font.color.rgb = COLOR_SECONDARY

    doc.add_paragraph()

    # ------------------ 4. PROJECTS SPECIFICATION ------------------
    add_h1("4. Featured Engineering Projects Catalog")
    doc.add_paragraph(
        "All projects presented in the portfolio are 100% authentic software prototypes and full-stack solutions built by Muruganantham G. "
        "Each project features custom UI cards, modal architecture breakdowns, live deployment and GitHub repository links:"
    )

    projects_list = [
        ("1. Real-Time Classroom Attendance AI", "AI & Computer Vision", "Python, OpenCV, Face Recognition, Django, PostgreSQL, HTML5, CSS3, JavaScript",
         "An automated classroom attendance system using facial recognition algorithms. Captures live video stream from classroom cameras, detects and identifies enrolled students in milliseconds, matches faces against registered biometric databases, and automatically marks attendance with date-time stamps, eliminating manual paper rolls and proxy attendance."),
        
        ("2. FarmDirect AI", "Agri-Tech & Marketplace", "Python, FastAPI, React, PostgreSQL, Scikit-learn, Tailwind CSS",
         "A direct-to-consumer agricultural commerce platform connecting regional farmers with local consumers and commercial buyers. Eliminates predatory middlemen, integrates AI price trend forecasting algorithms based on historical market data, and offers produce freshness verification and transparent order tracking."),
        
        ("3. Smart Campus R&D Portal", "Full Stack Web Application", "Next.js, TypeScript, Node.js, Express, MongoDB, Tailwind CSS",
         "A centralized research and development management portal for academic engineering institutions. Enables students, faculty, and research teams to submit project proposals, review peer papers, track patent application milestones, allocate laboratory equipment, and log collaborative research progress."),
        
        ("4. Advanced Hand Tracking AR", "Computer Vision & AR", "JavaScript, MediaPipe, Three.js, WebGL, Canvas API",
         "A browser-based interactive augmented reality playground tracking 21 3D hand landmarks in real time through webcam video feeds. Translates finger pinch, swipe, and palm gestures into real-time 3D object manipulation, virtual keyboard inputs, and interactive physics particle interactions with zero hardware sensors."),
        
        ("5. Smart AI Silage Testing", "Machine Learning & Quality QA", "Python, PyTorch, Convolutional Neural Networks (CNN), Flask, OpenCV",
         "An automated agricultural feed quality assessment tool leveraging computer vision to analyze agricultural silage fermentation and compaction. Evaluates moisture levels, mold distribution, and nutritional preservation from visual samples to give farmers actionable silage quality scores and safety alerts."),
        
        ("6. Student Growth AI", "Educational Analytics & ML", "Python, FastAPI, React, SQLite / PostgreSQL, Pandas, Chart.js",
         "An intelligent student progress tracking platform that models academic trajectories over multi-semester periods. Features predictive risk identification for struggling students, automated assignment submission monitoring, attendance correlation graphs, and personalized study module suggestions."),
        
        ("7. Personal 3D Developer Portfolio", "Creative Frontend & 3D WebGL", "Next.js 16, React 19, TypeScript, Three.js, Chibi3D, Tailwind CSS, Framer Motion",
         "The current production web application showcasing software engineering achievements through an interactive 3D mechanical keyboard, multi-route architecture, zero-scroll landing view, verified documentation viewers, and accessible theme toggles.")
    ]

    for title, cat, stack, desc in projects_list:
        add_h2(title)
        p = doc.add_paragraph()
        r_c = p.add_run(f"Category: {cat} | Tech Stack: {stack}\n")
        r_c.font.size = Pt(9.5)
        r_c.font.bold = True
        r_c.font.color.rgb = COLOR_SECONDARY
        p.add_run(desc).font.size = Pt(10)

    # ------------------ 5. CERTIFICATIONS DOCUMENTATION ------------------
    add_h1("5. Verified Certifications & Professional Credentials")
    doc.add_paragraph(
        "Muruganantham G has earned 10 verified technical certifications across Cloud Computing, Artificial Intelligence, Data Analytics, "
        "Human-Computer Interaction, DevOps, and Networking. Each credential is backed by official digital badge links or hosted high-resolution PDF certificates:"
    )

    cert_table = doc.add_table(rows=11, cols=4)
    cert_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cert_headers = ["Certification Title", "Issuing Body", "Key Competencies", "Verification Details"]
    for j, h in enumerate(cert_headers):
        cell = cert_table.rows[0].cells[j]
        set_cell_background(cell, HEX_HEADER)
        set_cell_margins(cell, 120, 120, 140, 140)
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)

    certs_data = [
        ("Explore Azure App Service", "Microsoft Learn", "Cloud app deployment, hosting, scaling, Azure App Service", "Completed June 3, 2026 • Verified PDF"),
        ("Enterprise Design Thinking", "IBM SkillsBuild", "Human-centered design, user research, empathy mapping, prototyping", "Issued Aug 05, 2026 • Credly ID: 5eyRbbTK"),
        ("Intro to Generative AI & Agents", "Microsoft", "GenAI fundamentals, autonomous agents, prompt engineering, ethics", "Completed May 30, 2026 • Verified PDF"),
        ("Intro to Computer Vision Concepts", "Microsoft", "Image processing, object detection, visual reasoning, CNNs", "Completed June 3, 2026 • Verified PDF"),
        ("Intro to Natural Language Processing", "Microsoft", "Text processing, semantic analysis, tokenization, language models", "Completed May 30, 2026 • Verified PDF"),
        ("Data Fundamentals", "IBM SkillsBuild", "Data management, analytics foundations, data architecture", "Issued Jul 09, 2026 • Credly ID: a8376efc..."),
        ("Human Computer Interaction", "NPTEL (IIT Madras & IIIT Delhi)", "Usability, user interface heuristics, cognitive psychology", "Elite 84% • Roll: NPTEL26CS70S250500380"),
        ("Learning Docker", "LinkedIn Learning", "Containers, Docker engine, microservices, container packaging", "Aug 07, 2026 • Certificate ID: 96e1662f..."),
        ("Introduction to SQL", "Simplilearn SkillUp", "SQL queries, relational joins, normalization, aggregations", "18th Nov 2025 • Code: 9422922"),
        ("Networking Essentials", "Cisco Networking Academy", "Network protocols, IP routing, subnetting, network security", "24 Dec 2025 • Cisco / Karpagam CE Verified")
    ]
    for i, row_data in enumerate(certs_data, start=1):
        row = cert_table.rows[i]
        bg = HEX_ALT if i % 2 == 1 else "FFFFFF"
        for j, text in enumerate(row_data):
            cell = row.cells[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(8.5)
            if j == 0:
                r.font.bold = True
            elif j == 1:
                r.font.color.rgb = COLOR_SECONDARY

    doc.add_paragraph()

    # ------------------ 6. EXPERIENCE & EDUCATION ------------------
    add_h1("6. Professional Experience & Academic Background")

    add_h2("6.1 Professional Experience")
    exp_items = [
        ("Full Stack Engineer (Intern)", "AEX Xenora Technologies", "Aug 2026 – Present | Remote (Coimbatore, Tamil Nadu)",
         "Selected for the Full Stack Web Development internship program (Ref: IOL - XEN244). Architecting responsive web interfaces with React, designing REST APIs with Python and FastAPI, querying relational databases, and participating in agile sprint reviews. Includes verified official 2-page Offer Letter PDF in application."),
        ("Web Developer (Intern)", "Yuva Intern by Henry Harvin", "Sep 2026 – Present | Remote (NSDC Accredited)",
         "Selected for the Full Stack Developer Fresher internship program accredited by Henry Harvin Education and the National Skill Development Corporation (NSDC). Developing accessible UI components, managing application state, and building modular frontend pipelines. Includes verified official Offer Letter PDF in application."),
        ("Industrial Trainee / Intern", "Texmo Industries", "Completed | Coimbatore, Tamil Nadu",
         "Gained hands-on industrial training at Texmo Industries, acquiring practical exposure to manufacturing workflows, quality assurance testing, technical troubleshooting, and workshop safety protocols.")
    ]
    for role, comp, dur, exp_desc in exp_items:
        add_h3(f"{role} — {comp}")
        p = doc.add_paragraph()
        r = p.add_run(f"{dur}\n")
        r.font.size = Pt(9.5)
        r.font.italic = True
        r.font.color.rgb = COLOR_MUTED
        p.add_run(exp_desc).font.size = Pt(10)

    add_h2("6.2 Academic Education")
    edu_items = [
        ("B.Tech in Information Technology", "Karpagam Institute of Technology, Coimbatore", "July 2025 – May 2028 (Pursuing)",
         "Specializing in Software Development, Cloud Computing, Database Architectures, and Applied AI. Earned NPTEL Elite certification (84%, 4 recommended credits) in Human-Computer Interaction."),
        ("Diploma in Electrical & Electronics Engineering (EEE)", "Sri Krishna Polytechnic College, Coimbatore", "June 2021 – June 2024",
         "Graduated with Distinction (85% aggregate score). Developed a rigorous foundation in circuit fundamentals, embedded systems, and engineering problem solving."),
        ("Secondary School Leaving Certificate (SSLC)", "Govt High School, Kurumbapalayam, Coimbatore", "Completed May 2021",
         "Graduated with 78% aggregate score, building early foundations in mathematics, physical sciences, and computer literacy.")
    ]
    for deg, inst, dur, edu_desc in edu_items:
        add_h3(f"{deg} — {inst}")
        p = doc.add_paragraph()
        r = p.add_run(f"{dur}\n")
        r.font.size = Pt(9.5)
        r.font.italic = True
        r.font.color.rgb = COLOR_MUTED
        p.add_run(edu_desc).font.size = Pt(10)

    # ------------------ 7. RECENT REFINEMENTS & OPTIMIZATIONS ------------------
    add_h1("7. Quality Enhancements & Customizations Implemented")
    doc.add_paragraph(
        "During iterative development and user alignment, several critical UX and technical refinements were executed:"
    )
    p = doc.add_paragraph()
    p.add_run("1. Zero-Scroll Home Page Optimization: ").font.bold = True
    p.add_run("Eliminated long, tedious home-page scroll-jacking. The Home page now fits cleanly within the viewport with the 3D keyboard and interactive controls.")
    p = doc.add_paragraph()
    p.add_run("2. Full-Width Animated Marquee Ribbon: ").font.bold = True
    p.add_run("Engineered a continuous, hardware-accelerated 60fps infinite marquee banner at the bottom of the hero displaying 'MURUGANANTHAM PORTFOLIO ✦ FULL STACK DEVELOPER ✦ AI ENTHUSIAST'.")
    p = doc.add_paragraph()
    p.add_run("3. Elimination of Card Scroll on Hover: ").font.bold = True
    p.add_run("Replaced jittery vertical auto-scrolling preview images on project cards with elegant static cover shots with subtle 1.04x scale zoom on hover.")
    p = doc.add_paragraph()
    p.add_run("4. Local High-Resolution SVG Assets: ").font.bold = True
    p.add_run("Fixed broken external CDN links for Tailwind CSS and AWS by generating optimized local vector assets (tailwind-cyan.svg and aws-orange.svg).")
    p = doc.add_paragraph()
    p.add_run("5. Contact Form Personalization: ").font.bold = True
    p.add_run("Updated the contact form Full Name field from a hardcoded developer name to a clean 'Your name' placeholder, and integrated direct telephone calling (+91 88706 42691).")
    p = doc.add_paragraph()
    p.add_run("6. Deduplication of Certifications: ").font.bold = True
    p.add_run("Removed the duplicate Certifications section from the Education page, creating a clear separation of concerns with its own dedicated /certifications route.")

    # ------------------ 8. GITHUB & DEPLOYMENT ------------------
    add_h1("8. GitHub Repository & Deployment Status")
    doc.add_paragraph(
        "The codebase is version-controlled with Git and hosted on GitHub under the candidate's personal account:"
    )
    doc.add_paragraph("• Repository URL: https://github.com/murugananthamg71-raji/MY-PERSONAL-3D-DEVELOPER-PORTFOLIO")
    doc.add_paragraph("• Default Branch: main (All working trees clean, fully synced)")
    doc.add_paragraph("• TypeScript Check: 100% Type-Safe (tsc --noEmit passed with Exit Code 0)")
    doc.add_paragraph("• Next.js Server: Next.js 16.2.2 Turbopack runtime active on port 3000")

    add_callout(
        "Deployment Readiness: The application is fully prepared for zero-configuration, continuous deployment on Vercel. "
        "All environment variables, dynamic routes, and static assets meet strict production optimization criteria.",
        italic=False
    )

    # ------------------ 9. CONCLUSION ------------------
    add_h1("9. Conclusion")
    doc.add_paragraph(
        "The Personal 3D Developer Portfolio successfully exemplifies modern web application development at the intersection of creative graphics and robust full-stack engineering. "
        "Through responsive architecture, accessible components, real-world project solutions, and verified credential transparency, Muruganantham G has constructed an impactful, "
        "memorable developer footprint that is ready for industry presentation and professional deployment."
    )

    output_path = "Muruganantham_G_Personal_3D_Developer_Portfolio_Project_Report.docx"
    doc.save(output_path)
    print(f"Report saved successfully to: {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_report()
